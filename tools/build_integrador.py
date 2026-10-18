"""Build the course deliverable from the supplied model and existing evidence.

Uses python-docx. Identity fields are deliberately unresolved until confirmed.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TITLE = 'Correspondências entre UML código Java e estado de objetos'

def build():
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        st = d.styles[name]
        st.font.name = 'Arial'
        st.font.size = Pt(12)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_after = Pt(6)
    d.styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for st in d.styles:
        for border in st.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for name in ['Heading 1', 'Heading 2', 'Heading 3']:
        d.styles[name].font.bold = True
        d.styles[name].paragraph_format.keep_with_next = True
        d.styles[name].paragraph_format.space_before = Pt(12)
    def para(t='', style=None):
        return d.add_paragraph(t, style)
    def centered(t, before=0, bold=False, style=None):
        p=para(t,style);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(before)
        if not style:
            for r in p.runs: r.bold=bold
        return p
    def head(t,level=1):
        p=d.add_heading(t,level);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        return p
    def page(): d.add_page_break()
    centered('ESCOLA [NOME A CONFIRMAR]',40,True)
    centered('CURSO TÉCNICO EM [NOME A CONFIRMAR]',0,True)
    centered('PROJETO INTEGRADOR',0,True)
    centered(TITLE.upper(),150,style='Title')
    centered('[CIDADE A CONFIRMAR]\nOUTUBRO DE 2026',205)
    page()
    centered('ALUNOS [NOMES A CONFIRMAR]',40,True)
    centered('PROFESSOR [NOME A CONFIRMAR]',0,True)
    centered(TITLE.upper(),150,style='Title')
    centered('[CIDADE A CONFIRMAR]\nOUTUBRO DE 2026',225)
    page()
    centered('SUMÁRIO',0,True)
    toc = [
        ('1 INTRODUÇÃO','intro'),('1.1 OBJETIVOS','obj'),
        ('1.1.1 Objetivo Geral','geral'),('1.1.2 Objetivos Específicos','especificos'),
        ('1.2 METODOLOGIA','metodo'),('2 REFERENCIAL TEÓRICO','teoria'),
        ('3 CONSIDERAÇÕES FINAIS','fim'),('4 REFERÊNCIAS','refs'),
        ('APÊNDICE A Verificação técnica do demonstrador','apendice'),('ANEXOS','anexos')]
    page_numbers = dict(zip([mark for _,mark in toc], [4,4,4,4,5,6,8,9,10,11]))
    for label,mark in toc:
        p=para(label+'\t');p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGEREF '+mark+' \\h')
        r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=str(page_numbers[mark]);r.append(t);f.append(r);p._p.append(f)
    marks={label:mark for label,mark in toc}
    count=0
    def marked(t,level=1):
        nonlocal count
        p=head(t,level);count+=1
        b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(count));b.set(qn('w:name'),marks[t])
        e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(count));p._p.insert(0,b);p._p.append(e)
    page();marked('1 INTRODUÇÃO')
    para('Este Projeto Integrador apresenta um demonstrador que relaciona a estrutura de uma classe, o código Java e o estado de seus objetos. O exemplo utiliza contas bancárias para mostrar que uma classe pode originar instâncias com valores independentes e que duas referências podem apontar para o mesmo objeto. O resultado desenvolvido é um recurso técnico de visualização, acompanhado de verificações reproduzíveis.')
    para('O problema abordado é como tornar observável a correspondência entre um membro de classe, sua implementação e a mudança de estado do objeto que recebe uma chamada. A pergunta deste recorte técnico é: como implementar essas correspondências e verificar seu comportamento nos cenários definidos? Essa pergunta pode ser examinada por análise documental e testes de software. O efeito sobre a aprendizagem exige outra avaliação, com participantes, que não foi realizada.')
    para('A proposta dialoga com antecedentes como o BlueJ, que já combina representações de classes e interação com objetos (Kölling et al., 2003). Portanto, não se reivindica a primeira integração entre UML e execução. O interesse do demonstrador está na explicitação das correspondências, na comparação entre duas condições visuais e na preservação das evidências de funcionamento.')
    marked('1.1 OBJETIVOS',2)
    marked('1.1.1 Objetivo Geral',3)
    para('Desenvolver e verificar um demonstrador de correspondências entre uma representação estrutural UML, o código Java e o estado de objetos em cenários delimitados de orientação a objetos.')
    marked('1.1.2 Objetivos Específicos',3)
    para('Identificar fundamentos e recursos de ferramentas relacionadas para justificar as correspondências propostas entre classe, código e instâncias.')
    para('Implementar e verificar os cenários de instâncias independentes, referências compartilhadas e mudanças de estado, registrando os resultados técnicos e suas limitações.')
    marked('1.2 METODOLOGIA',2)
    para('O trabalho combina pesquisa bibliográfica exploratória, análise documental de ferramentas e desenvolvimento de um protótipo com verificação técnica. A articulação entre construir e avaliar artefatos se aproxima da perspectiva de Design Science Research apresentada por Hevner et al. (2004). Este recorte não corresponde à conclusão de uma avaliação educacional do artefato.')
    para('A etapa bibliográfica consultou publicações e registros institucionais sobre ambientes de aprendizagem de Java e visualização de programas. O levantamento foi exploratório: algumas fontes tiveram seções consultadas, outras apenas resumos e metadados. A busca não foi uma revisão sistemática. O corpus e os limites de leitura estão registrados na matriz de evidências do projeto.')
    para('Na implementação, utilizou-se a classe BankAccount, com saldo decimal, construtor e métodos de depósito, retirada e consulta. Um executor mantém objetos reais em Java e registra cópias imutáveis dos estados antes e depois das operações. Marcadores delimitam as regiões de fonte associadas aos membros. A interface oferece as condições linked, com destaques de correspondência, e plain, sem esses destaques, utilizando o mesmo executor.')
    para('A verificação empregou valores conhecidos e cenários determinísticos, sem participantes. Foram compilados o exemplo executável e um esqueleto gerado separadamente. Os testes conferiram identidade de objetos, independência entre instâncias, referências compartilhadas, valores decimais, rejeição de operações inválidas, equivalência do executor e registro de respostas. Uma execução adicional acionou a interface em display virtual e conferiu a exportação local.')
    para('As evidências foram interpretadas segundo o comportamento esperado de cada cenário. Passar em um teste sustenta apenas a afirmação correspondente ao cenário exercitado; não demonstra correção para todos os programas, facilidade de uso, aprendizagem ou potencial de venda. Fontes, versões, logs e hashes estão preservados no repositório público do projeto. Os procedimentos e resultados estão resumidos no Apêndice A.')
    page();marked('2 REFERENCIAL TEÓRICO')
    head('2.1 Classes objetos e representações relacionadas',2)
    para('Uma classe descreve a estrutura e os comportamentos de seus objetos. O estado pertence a cada instância; uma referência permite acessar um objeto e não deve ser confundida com ele. No exemplo adotado, criar duas contas com saldos distintos permite observar a independência dos estados. Criar um segundo nome para a mesma conta permite observar referências compartilhadas. Essas situações orientam a escolha do exemplo e os testes do demonstrador.')
    para('O BlueJ oferece um antecedente relevante para essa aproximação entre estrutura e interação. Kölling et al. (2003) descrevem um ambiente voltado à aprendizagem de orientação a objetos, com representação de classes e interação direta com objetos. Esse antecedente impede atribuir novidade ampla à simples presença de um diagrama e de execução em uma mesma ferramenta. A comparação aqui é documental, sem ensaio instalado de desempenho ou usabilidade entre produtos.')
    head('2.2 Múltiplas representações e atividades de compreensão',2)
    para('Ainsworth (2006) propõe o framework DeFT para considerar o desenho, as funções e as tarefas cognitivas envolvidas no uso de múltiplas representações. Sua contribuição para este projeto é orientar a relação entre as vistas: o cartão estrutural indica os membros, o código mostra a implementação e os estados mostram os valores dos objetos. A seleção de um membro e a execução de um método evidenciam correspondências entre essas informações.')
    para('A disponibilização simultânea de informações não comprova que o estudante compreendeu a relação entre elas. O demonstrador inclui campos para prever uma mudança e explicar o resultado observado. A síntese de estudos de visualização de algoritmos de Hundhausen, Douglas e Stasko (2002) oferece fundamento para considerar as atividades realizadas com a visualização. A aplicação desse argumento ao exemplo de objetos é uma decisão de projeto que ainda requer avaliação educacional própria.')
    head('2.3 Visualização dinâmica e limites da contribuição',2)
    para('Moreno et al. (2004) apresentam o Jeliot 3 como ferramenta de visualização de programas. Gestwicki e Jayaraman (2004) descrevem o JIVE, relacionado à visualização da execução em Java. Esses antecedentes reforçam que a representação visual do comportamento de programas já integra uma linha de desenvolvimento existente. O demonstrador deste trabalho não substitui a análise dessas ferramentas nem prova ausência de funcionalidades semelhantes nelas.')
    para('O recurso implementado é restrito a um exemplo de conta bancária. Na condição linked, a seleção de um membro destaca sua região de código e a operação evidencia as referências afetadas. Na condição plain, os estados e as operações permanecem os mesmos, sem o apoio visual adicional. Essa separação torna possível planejar uma comparação futura, embora a existência das duas condições não constitua, por si, um experimento com estudantes.')
    head('2.4 Construção avaliação e comunicação do produto',2)
    para('Hevner et al. (2004) articulam a construção e a avaliação de artefatos em Sistemas de Informação. Para o presente trabalho, essa perspectiva orienta a necessidade de declarar qual problema o recurso aborda e quais evidências foram obtidas. O protótipo e seus testes sustentam um resultado técnico delimitado. Afirmar melhoria de aprendizagem, utilidade percebida ou interesse comercial demandaria evidências adicionais.')
    para('Uma comunicação compatível com o resultado demonstrado pode apresentar o recurso como uma forma de visualizar relações entre classe, código e mudanças de estado em objetos Java. Não seria justificável anunciar aprendizagem comprovada, geração automática de lógica de negócio, suporte a Java arbitrário ou integração inédita. A pesquisa de demanda e a avaliação de disposição a pagar permanecem etapas distintas da verificação funcional.')
    page();marked('3 CONSIDERAÇÕES FINAIS')
    para('O objetivo técnico foi desenvolvido mediante um demonstrador que relaciona uma representação estrutural UML, regiões de fonte Java e estados de objetos reais. A primeira verificação preservada apresentou 28 verificações Java e cinco testes Python aprovados. Em uma execução posterior, nove verificações de interface e exportação também foram aprovadas. O Apêndice A identifica as versões e as execuções, permitindo conferir o alcance desses resultados.')
    para('Os cenários verificados mostram que depositar em uma conta não altera uma instância independente e que referências compartilhadas observam a mudança do mesmo objeto. Também foram conferidos registros de estado e a equivalência dos comportamentos entre as duas condições visuais. Isso responde à pergunta técnica dentro dos exemplos exercitados, sem demonstrar efeito sobre a aprendizagem.')
    para('A análise bibliográfica identificou antecedentes que impedem tratar a simples união de diagrama e objetos como contribuição inédita. O resultado deve ser apresentado como um demonstrador restrito, com correspondências explícitas e evidências rastreáveis. Os objetivos de identificar fundamentos e implementar cenários foram operacionalizados; a revisão crítica aprofundada e a avaliação educacional do programa de pesquisa mais amplo permanecem em andamento.')
    para('Entre as limitações estão o uso de uma única classe, o mapeamento baseado em marcadores e a ausência de avaliação humana de usabilidade ou aprendizagem. O protótipo não é um editor UML geral e não executa qualquer programa Java. Como continuidade, propõe-se revisar os instrumentos, confirmar as condições institucionais e realizar um estudo adequado antes de afirmar ganho educacional. A demanda comercial também deve ser investigada por procedimentos próprios.')
    page();marked('4 REFERÊNCIAS')
    manuscript=(ROOT/'paper/manuscrito.md').read_text()
    refs=manuscript.split('## Referências\n',1)[1]
    for item in refs.strip().split('\n\n'):
        item=item.split(' O registro Crossref')[0]
        if item.startswith('GESTWICKI'): item=item.replace('JIVE: Java interactive visualization environment.', 'JIVE.')
        p=para(item);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    page();marked('APÊNDICE A Verificação técnica do demonstrador')
    para('O demonstrador está em prototype/ no repositório aloisiocosta-professor/trama-ide-web. Os testes usam dados de exemplo, sem registros de pessoas participantes. Os números abaixo representam verificações de software, não quantidade de estudantes.')
    head('Cenários e resultados',2)
    para('A execução inicial 39083293325 do GitHub Actions corresponde ao commit 5cf946e5a77439c3c41759536c006d6d8a2e3810. Nela passaram 28 verificações Java e cinco testes Python do gerador. A execução 39104951676 corresponde ao commit 89a84aa3132eb6c369a2dfe7ff2057716fe4b0df e acrescenta nove verificações gráficas de janela, controles, legibilidade, respostas e exportação. Os dois conjuntos são evidências de versões diferentes, preservadas separadamente.')
    para('O cenário principal cria duas contas, seleciona uma delas, registra uma previsão e executa um depósito de 25 sobre um saldo inicial de 100. O saldo dessa conta passa a 125, enquanto a outra instância mantém seu saldo de 20. Os testes também conferem alias e rejeições; o cenário principal não representa todos os comportamentos possíveis.')
    head('Como conferir as evidências',2)
    para('O comando python tools/build_research.py requer JDK 17, Python e xvfb-run para as verificações gráficas em Linux. Os logs, resultados e imagens das execuções citadas estão em docs/research/evidence/ci-5cf946/ e docs/research/evidence/ci-89a84a/. O README do protótipo documenta execução interativa e limitações. A compilação foi verificada no ambiente remoto descrito nos logs.')
    para('Repositório: https://github.com/aloisiocosta-professor/trama-ide-web/tree/docs/research-bootstrap . Execuções: https://github.com/aloisiocosta-professor/trama-ide-web/actions/runs/39083293325 e https://github.com/aloisiocosta-professor/trama-ide-web/actions/runs/39104951676 .')
    para('Houve assistência de inteligência artificial na organização textual e no desenvolvimento dos materiais. Cabe aos autores conferir fontes, compreender o funcionamento, revisar o texto e assumir responsabilidade pela entrega. A forma de declaração deve observar as regras da escola.')
    page();marked('ANEXOS')
    para('Não foram incorporados documentos de terceiros ao corpo desta versão. O modelo institucional enviado orientou a estrutura e a formatação. As fontes bibliográficas utilizadas estão identificadas nas referências; seus textos integrais não foram reproduzidos.')
    d.core_properties.title=TITLE
    d.core_properties.author=''
    d.core_properties.subject='Projeto Integrador'
    d.core_properties.comments=''
    out=ROOT/'paper/projeto-integrador-versao-revisao.docx'
    d.save(out)
    return out

if __name__=='__main__': print(build())
